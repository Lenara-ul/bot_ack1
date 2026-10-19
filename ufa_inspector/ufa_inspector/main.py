"""Прототип: автоматизация оформления актов по фото (ЦОБ г. Уфа).
Запуск: uvicorn main:app --reload   →   http://127.0.0.1:8000
"""
import csv, io, sqlite3, uuid
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from PIL import Image
from pydantic import BaseModel

BASE = Path(__file__).parent
UPLOADS = BASE / "uploads"
UPLOADS.mkdir(exist_ok=True)
DB_PATH = BASE / "acts.db"

# ---------------------------------------------------------------------------
# Справочник нарушений. ВАЖНО: номера статей КоАП РБ и пунктов Правил
# благоустройства г. Уфы ниже — ДЕМО-ЗАГЛУШКИ. Перед использованием их должен
# проверить и заполнить юрист (актуальные редакции документов).
# ---------------------------------------------------------------------------
VIOLATIONS = [
    {"code": "lawn_parking", "title": "Парковка на газоне",
     "koap": "ст. ___ КоАП РБ [УТОЧНИТЬ]", "rules": "п. ___ Правил благоустройства г. Уфы [УТОЧНИТЬ]"},
    {"code": "trash_entrance", "title": "Мусор у подъезда",
     "koap": "ст. ___ КоАП РБ [УТОЧНИТЬ]", "rules": "п. ___ Правил благоустройства г. Уфы [УТОЧНИТЬ]"},
    {"code": "blocked_passage", "title": "Перекрытие проезда",
     "koap": "ст. ___ КоАП РБ [УТОЧНИТЬ]", "rules": "п. ___ Правил благоустройства г. Уфы [УТОЧНИТЬ]"},
    {"code": "playground_parking", "title": "Парковка на детской/спортивной площадке",
     "koap": "ст. ___ КоАП РБ [УТОЧНИТЬ]", "rules": "п. ___ Правил благоустройства г. Уфы [УТОЧНИТЬ]"},
]
V_BY_CODE = {v["code"]: v for v in VIOLATIONS}

# ---------------------------------------------------------------------------
# БД (SQLite, один файл acts.db)
# ---------------------------------------------------------------------------
def db():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c

with db() as c:
    c.execute("""CREATE TABLE IF NOT EXISTS acts(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT, time TEXT, address TEXT, plate TEXT,
        violation_code TEXT, violation_title TEXT, norms TEXT,
        status TEXT, inspector_name TEXT, inspector_position TEXT,
        photo TEXT, created_at TEXT)""")

# ---------------------------------------------------------------------------
# Загрузка файла: EXIF + Mock-распознавание
# ---------------------------------------------------------------------------
def read_exif(img):
    """Возвращает (datetime|None, (lat, lon)|None) из EXIF."""
    dt, gps = None, None
    try:
        ex = img.getexif()
        raw = ex.get_ifd(0x8769).get(36867) or ex.get(306)  # DateTimeOriginal / DateTime
        if raw:
            dt = datetime.strptime(str(raw), "%Y:%m:%d %H:%M:%S")
        g = ex.get_ifd(0x8825)
        if g and 2 in g and 4 in g:
            deg = lambda v, ref: (-1 if ref in ("S", "W") else 1) * (
                float(v[0]) + float(v[1]) / 60 + float(v[2]) / 3600)
            gps = (round(deg(g[2], g.get(1)), 6), round(deg(g[4], g.get(3)), 6))
    except Exception:
        pass
    return dt, gps

def mock_ocr(path):
    # Здесь будет интеграция с OCR API (например, Yandex Vision / Tesseract)
    return {"plate": "А123ВС102", "datetime": datetime(2026, 10, 19, 9, 30)}

app = FastAPI(title="Акты ЦОБ Уфа")

@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(400, "Нужен файл-изображение (JPG/PNG)")
    name = f"{uuid.uuid4().hex}{Path(file.filename or 'x.jpg').suffix.lower() or '.jpg'}"
    path = UPLOADS / name
    path.write_bytes(await file.read())
    try:
        exif_dt, gps = read_exif(Image.open(path))
    except Exception:
        path.unlink(missing_ok=True)
        raise HTTPException(400, "Не удалось открыть изображение")
    ocr = mock_ocr(path)
    # Приоритет даты/времени: EXIF → текст на фото (mock) → ручной ввод
    if exif_dt:
        dt, src = exif_dt, "EXIF метаданных фото"
    elif ocr.get("datetime"):
        dt, src = ocr["datetime"], "текста на фото (демо-данные)"
    else:
        dt, src = None, None
    return {
        "photo": name,
        "plate": ocr["plate"],
        "date": dt.strftime("%Y-%m-%d") if dt else "",
        "time": dt.strftime("%H:%M") if dt else "",
        "datetime_source": src,
        "gps": {"lat": gps[0], "lon": gps[1]} if gps else None,
    }

# ---------------------------------------------------------------------------
# Сохранение в БД
# ---------------------------------------------------------------------------
class ActIn(BaseModel):
    id: int | None = None
    date: str = ""
    time: str = ""
    address: str = ""
    plate: str = ""
    violation_code: str = ""
    inspector_name: str = ""
    inspector_position: str = ""
    photo: str = ""
    status: str = "черновик"

@app.get("/api/violations")
def violations():
    return VIOLATIONS

@app.post("/api/acts")
def save_act(a: ActIn):
    if a.status not in ("черновик", "оформлен"):
        raise HTTPException(400, "Неверный статус")
    if a.status == "оформлен":  # для окончательного акта все поля обязательны
        missing = [n for n, v in [("дата", a.date), ("время", a.time), ("адрес", a.address),
                                  ("госномер", a.plate), ("нарушение", a.violation_code),
                                  ("ФИО инспектора", a.inspector_name),
                                  ("должность", a.inspector_position)] if not v.strip()]
        if missing:
            raise HTTPException(400, "Заполните поля: " + ", ".join(missing))
    v = V_BY_CODE.get(a.violation_code, {})
    norms = f"{v.get('koap', '')}; {v.get('rules', '')}" if v else ""
    vals = (a.date, a.time, a.address.strip(), a.plate.strip().upper(), a.violation_code,
            v.get("title", ""), norms, a.status, a.inspector_name.strip(),
            a.inspector_position.strip(), a.photo)
    with db() as c:
        if a.id:
            c.execute("""UPDATE acts SET date=?,time=?,address=?,plate=?,violation_code=?,
                violation_title=?,norms=?,status=?,inspector_name=?,inspector_position=?,photo=?
                WHERE id=?""", vals + (a.id,))
            act_id = a.id
        else:
            act_id = c.execute("""INSERT INTO acts(date,time,address,plate,violation_code,
                violation_title,norms,status,inspector_name,inspector_position,photo,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""", vals + (datetime.now().isoformat(),)).lastrowid
    return {"id": act_id, "status": a.status}

@app.get("/api/acts")
def list_acts(limit: int | None = None):
    sql = "SELECT * FROM acts ORDER BY date DESC, time DESC, id DESC"
    with db() as c:
        rows = c.execute(sql + (" LIMIT ?" if limit else ""), (limit,) if limit else ()).fetchall()
    return [dict(r) for r in rows]

def get_act(act_id: int):
    with db() as c:
        r = c.execute("SELECT * FROM acts WHERE id=?", (act_id,)).fetchone()
    if not r:
        raise HTTPException(404, "Акт не найден")
    return dict(r)

@app.get("/api/acts/{act_id}")
def one_act(act_id: int):
    return get_act(act_id)

# ---------------------------------------------------------------------------
# Формирование акта (DOCX по шаблону, собранному в коде)
# ---------------------------------------------------------------------------
def build_docx(a: dict) -> bytes:
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)
    for text, bold in [("Центр общественной безопасности города Уфа", True),
                       (f"АКТ № {a['id']} фиксации нарушения", True)]:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(text).bold = bold
    y, m, dd = (a["date"] or "--").split("-") if a["date"] else ("", "", "")
    rows = [("Дата и время фиксации", f"{dd}.{m}.{y} {a['time']}"),
            ("Место фиксации (адрес)", a["address"]),
            ("Государственный номер ТС", a["plate"]),
            ("Выявленное нарушение", a["violation_title"]),
            ("Нормы", a["norms"]),
            ("Инспектор (ФИО)", a["inspector_name"]),
            ("Должность", a["inspector_position"])]
    t = d.add_table(rows=0, cols=2, style="Table Grid")
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text = k, v or ""
        cells[0].paragraphs[0].runs[0].bold = True
    if a.get("photo") and (UPLOADS / a["photo"]).exists():
        d.add_paragraph("\nФотофиксация:")
        try:
            d.add_picture(str(UPLOADS / a["photo"]), width=Cm(10))
        except Exception:
            d.add_paragraph("[фото не удалось вставить]")
    d.add_paragraph("\n")
    # Рамки для подписи и печати (пустые поля)
    s = d.add_table(rows=1, cols=2, style="Table Grid")
    s.rows[0].height = Cm(3)
    s.rows[0].cells[0].text = f"Подпись инспектора:\n\n______________ / {a['inspector_name']}"
    s.rows[0].cells[1].text = "Место для печати (М.П.)"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

@app.get("/api/acts/{act_id}/docx")
def act_docx(act_id: int):
    data = build_docx(get_act(act_id))
    return Response(data, headers={"Content-Disposition": f'attachment; filename="akt_{act_id}.docx"'},
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

# ---------------------------------------------------------------------------
# Выгрузка CSV (';' и BOM — чтобы корректно открывалось в Excel)
# ---------------------------------------------------------------------------
@app.get("/api/acts.csv")
def acts_csv():
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Дата", "Время", "Адрес", "Номер авто", "Тип нарушения",
                "Статьи/пункты", "Статус", "ФИО инспектора"])
    for r in list_acts():
        w.writerow([r["date"], r["time"], r["address"], r["plate"], r["violation_title"],
                    r["norms"], r["status"], r["inspector_name"]])
    return Response("\ufeff" + out.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="reestr_aktov.csv"'})

app.mount("/uploads", StaticFiles(directory=UPLOADS), name="uploads")
app.mount("/", StaticFiles(directory=BASE / "static", html=True), name="static")
